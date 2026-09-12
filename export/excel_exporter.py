"""
InvoiceIQ - Excel Exporter

Converts extracted Invoice objects into a structured
Excel workbook.

Workbook structure:

    Invoices
    Items
    Taxes
    Validation

The design intentionally separates invoice-level,
item-level, and tax-level information so that
multi-item invoices are represented correctly.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from extraction.invoice_extractor import Invoice
from extraction.validators import validate_invoice


# ============================================================
# CONFIGURATION
# ============================================================

DEFAULT_OUTPUT = Path("output") / "InvoiceIQ.xlsx"


# ============================================================
# HEADERS
# ============================================================

INVOICE_HEADERS = [
    "Invoice Number",
    "Invoice Type",
    "Invoice Date",
    "Due Date",
    "Seller",
    "Seller GSTIN",
    "Buyer",
    "Buyer GSTIN",
    "Total Tax",
    "Round Off",
    "Total Amount",
    "Received Amount",
    "Validation Status",
]

ITEM_HEADERS = [
    "Invoice Number",
    "Item No",
    "Material Name",
    "HSN",
    "Quantity",
    "Unit",
    "Rate",
    "Taxable Value",
]

TAX_HEADERS = [
    "Invoice Number",
    "Tax Type",
    "Rate (%)",
    "Amount",
]

VALIDATION_HEADERS = [
    "Invoice Number",
    "Status",
    "Valid",
    "Errors",
    "Warnings",
]


# ============================================================
# STYLING
# ============================================================


def style_header(ws) -> None:
    """
    Apply consistent formatting to a worksheet header.
    """

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def auto_size_columns(ws) -> None:
    """
    Automatically adjust column widths based on content.
    """

    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is None:
                continue

            length = len(
                str(cell.value)
            )

            max_length = max(
                max_length,
                length,
            )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            50,
        )


def prepare_sheet(ws) -> None:
    """
    Apply common worksheet formatting.
    """

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = ws.dimensions

    auto_size_columns(ws)


# ============================================================
# INVOICES SHEET
# ============================================================


def write_invoices_sheet(
    workbook: Workbook,
    invoices: list[Invoice],
) -> None:
    """
    Write invoice-level information.
    """

    ws = workbook.create_sheet(
        "Invoices"
    )

    ws.append(
        INVOICE_HEADERS
    )

    for invoice in invoices:

        validation = validate_invoice(
            invoice
        )

        ws.append(
            [
                invoice.invoice_number,
                invoice.invoice_type,
                invoice.invoice_date,
                invoice.due_date,
                invoice.seller,
                invoice.seller_gstin,
                invoice.buyer,
                invoice.buyer_gstin,
                invoice.total_tax,
                invoice.round_off,
                invoice.total_amount,
                invoice.received_amount,
                validation.status,
            ]
        )

    style_header(ws)
    prepare_sheet(ws)


# ============================================================
# ITEMS SHEET
# ============================================================


def write_items_sheet(
    workbook: Workbook,
    invoices: list[Invoice],
) -> None:
    """
    Write every invoice item as a separate row.

    This is important for multi-item invoices.
    """

    ws = workbook.create_sheet(
        "Items"
    )

    ws.append(
        ITEM_HEADERS
    )

    for invoice in invoices:

        for item_number, item in enumerate(
            invoice.items,
            start=1,
        ):

            ws.append(
                [
                    invoice.invoice_number,
                    item_number,
                    item.material_name,
                    item.hsn,
                    item.quantity,
                    item.unit,
                    item.rate,
                    item.taxable_value,
                ]
            )

    style_header(ws)
    prepare_sheet(ws)


# ============================================================
# TAXES SHEET
# ============================================================


def write_taxes_sheet(
    workbook: Workbook,
    invoices: list[Invoice],
) -> None:
    """
    Write every invoice tax as a separate row.
    """

    ws = workbook.create_sheet(
        "Taxes"
    )

    ws.append(
        TAX_HEADERS
    )

    for invoice in invoices:

        for tax in invoice.taxes:

            ws.append(
                [
                    invoice.invoice_number,
                    tax.type,
                    tax.rate,
                    tax.amount,
                ]
            )

    style_header(ws)
    prepare_sheet(ws)


# ============================================================
# VALIDATION SHEET
# ============================================================


def write_validation_sheet(
    workbook: Workbook,
    invoices: list[Invoice],
) -> None:
    """
    Write validation results separately so that
    accounting users can quickly identify problems.
    """

    ws = workbook.create_sheet(
        "Validation"
    )

    ws.append(
        VALIDATION_HEADERS
    )

    for invoice in invoices:

        result = validate_invoice(
            invoice
        )

        ws.append(
            [
                invoice.invoice_number,
                result.status,
                result.is_valid,
                "\n".join(result.errors),
                "\n".join(result.warnings),
            ]
        )

    style_header(ws)
    prepare_sheet(ws)


# ============================================================
# EXPORT MULTIPLE INVOICES
# ============================================================


def export_invoices_excel(
    invoices: list[Invoice],
    output_path: str | Path = DEFAULT_OUTPUT,
) -> Path:
    """
    Export multiple Invoice objects to an Excel workbook.

    Sheets created:

        1. Invoices
        2. Items
        3. Taxes
        4. Validation
    """

    output_path = Path(
        output_path
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    # Remove the default empty worksheet.
    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )

    write_invoices_sheet(
        workbook,
        invoices,
    )

    write_items_sheet(
        workbook,
        invoices,
    )

    write_taxes_sheet(
        workbook,
        invoices,
    )

    write_validation_sheet(
        workbook,
        invoices,
    )

    workbook.save(
        output_path
    )

    return output_path


# ============================================================
# EXPORT ONE INVOICE
# ============================================================


def export_invoice_excel(
    invoice: Invoice,
    output_path: str | Path = DEFAULT_OUTPUT,
) -> Path:
    """
    Convenience function for exporting one invoice.
    """

    return export_invoices_excel(
        [invoice],
        output_path,
    )


# ============================================================
# COMMAND LINE TEST
# ============================================================


if __name__ == "__main__":

    import sys

    from extraction.invoice_extractor import (
        extract_invoice,
    )

    if len(sys.argv) < 2:

        print(
            "Usage:"
        )

        print(
            "python -m export.excel_exporter "
            "<invoice.pdf> [output.xlsx]"
        )

        sys.exit(1)

    pdf_path = Path(
        sys.argv[1]
    )

    if len(sys.argv) >= 3:

        output_path = Path(
            sys.argv[2]
        )

    else:

        output_path = DEFAULT_OUTPUT

    print(
        f"Processing: {pdf_path.name}"
    )

    invoice = extract_invoice(
        pdf_path
    )

    validation = validate_invoice(
        invoice
    )

    print(
        f"Validation: {validation.status}"
    )

    if validation.errors:

        print("\nErrors:")

        for error in validation.errors:

            print(
                f"  ❌ {error}"
            )

    if validation.warnings:

        print("\nWarnings:")

        for warning in validation.warnings:

            print(
                f"  ⚠️ {warning}"
            )

    output = export_invoice_excel(
        invoice,
        output_path,
    )

    print(
        f"\n✅ Excel exported:"
    )

    print(
        output
    )
