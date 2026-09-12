"""
Tests for the InvoiceIQ Excel exporter.
"""

from pathlib import Path

from openpyxl import load_workbook

from extraction.invoice_extractor import extract_invoice
from export.excel_exporter import export_invoice_excel


# ============================================================
# Paths
# ============================================================

PROJECT_ROOT = (
    Path(__file__).resolve().parent.parent
)

SAMPLES_DIR = (
    PROJECT_ROOT / "samples"
)

PURCHASE_03 = (
    SAMPLES_DIR / "Purchase Invoice 03.pdf"
)


# ============================================================
# Excel Export Test
# ============================================================


def test_purchase_invoice_excel_export(
    tmp_path,
):

    invoice = extract_invoice(
        PURCHASE_03
    )

    output_file = (
        tmp_path / "InvoiceIQ.xlsx"
    )

    result = export_invoice_excel(
        invoice,
        output_file,
    )

    # --------------------------------------------------------
    # File exists
    # --------------------------------------------------------

    assert result.exists()

    # --------------------------------------------------------
    # Open workbook
    # --------------------------------------------------------

    workbook = load_workbook(
        result
    )

    # --------------------------------------------------------
    # Sheets
    # --------------------------------------------------------

    assert set(
        workbook.sheetnames
    ) == {
        "Invoices",
        "Items",
        "Taxes",
        "Validation",
    }

    # --------------------------------------------------------
    # Invoices sheet
    # --------------------------------------------------------

    invoices = workbook[
        "Invoices"
    ]

    assert invoices.max_row == 2

    assert (
        invoices["A2"].value
        == "KGSDF/YGEPL/26-27/03"
    )

    assert (
        invoices["B2"].value
        == "PURCHASE"
    )

    assert (
        invoices["K2"].value
        == 158132.0
    )

    assert (
        invoices["M2"].value
        == "VALID"
    )

    # --------------------------------------------------------
    # Items sheet
    # --------------------------------------------------------

    items = workbook[
        "Items"
    ]

    # Header + 2 items
    assert items.max_row == 3

    assert (
        items["B2"].value
        == 1
    )

    assert (
        items["C2"].value
        == "SAW DUST PELLETS PREMIUM"
    )

    assert (
        items["B3"].value
        == 2
    )

    assert (
        items["C3"].value
        == "SAW DUST BRIQUETTE PREMIUM"
    )

    # --------------------------------------------------------
    # Taxes sheet
    # --------------------------------------------------------

    taxes = workbook[
        "Taxes"
    ]

    # Header + CGST + SGST
    assert taxes.max_row == 3

    assert (
        taxes["B2"].value
        == "CGST"
    )

    assert (
        taxes["D2"].value
        == 3765.05
    )

    assert (
        taxes["B3"].value
        == "SGST"
    )

    assert (
        taxes["D3"].value
        == 3765.05
    )

    # --------------------------------------------------------
    # Validation sheet
    # --------------------------------------------------------

    validation = workbook[
        "Validation"
    ]

    assert validation.max_row == 2

    assert (
        validation["B2"].value
        == "VALID"
    )

    assert (
        validation["C2"].value
        is True
    )

    workbook.close()
